"""Enterprise ingestion and cross-format assurance; no provider network."""
import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from unittest.mock import Mock
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from starlette.datastructures import UploadFile
from app.main import app
from app.ingestion.bundle_reader import ingest_uploads
from app.ingestion.document_reader import read_document
from app.ingestion.field_mapper import convert, map_table
from app.ingestion.models import IngestionError
from app.ingestion.spreadsheet_reader import read_xlsx
from app.ingestion.csv_reader import read_csv
from app.services.sample_loader import load_sample_evidence, SAMPLE_DIRECTORY
from app.services import assessment_orchestrator
from test_ai_assessment import extraction_result

ENTERPRISE = SAMPLE_DIRECTORY.parent / 'assesspro_enterprise'
PATH = '/api/v1/ai/assess-files'

def payload(enterprise=True):
    names = {'status_report': 'weekly_status_report.docx' if enterprise else 'weekly_status_report.md',
             'raid_log': 'raid_log.xlsx' if enterprise else 'raid_log.json',
             'milestone_plan': 'milestone_plan.xlsx' if enterprise else 'milestone_plan.json',
             'quality_summary': 'defect_summary.xlsx' if enterprise else 'defect_summary.json'}
    directory = ENTERPRISE if enterprise else SAMPLE_DIRECTORY
    return {field: (name, (directory / name).read_bytes()) for field, name in names.items()}

def uploads(files):
    return {field: UploadFile(filename=name, file=BytesIO(data)) for field, (name, data) in files.items()}

def pdf(text='', encrypted=False):
    writer = PdfWriter(); page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject(); stream.set_data(('BT /F1 12 Tf 30 750 Td ('+text+') Tj ET').encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    if encrypted: writer.encrypt('test-password')
    out=BytesIO(); writer.write(out); return out.getvalue()

def test_docx_body_table_order_and_headers():
    result = read_document('report.docx', '.docx', (ENTERPRISE / 'weekly_status_report.docx').read_bytes())
    assert result.file_type == 'docx' and result.character_count == len(result.extracted_text)
    assert 'Project: AssessPro Modernization & Feature Expansion' in result.extracted_text
    assert result.extracted_text.index('Project:') < result.extracted_text.index('Delivery update')
    assert '11 Severity-2 scoring defects remain open.' in result.extracted_text

def test_text_pdf():
    doc = read_document('report.pdf', '.pdf', pdf('Development is on track. Certification is still pending and must be reviewed.'))
    assert 'Certification is still pending' in doc.extracted_text and doc.page_count == 1

@pytest.mark.parametrize('data,reason', [(pdf(), 'pdf_text'), (pdf(encrypted=True), 'encrypted'), (b'broken', 'malformed')])
def test_pdf_rejections(data, reason):
    with pytest.raises(IngestionError) as exc: read_document('report.pdf', '.pdf', data)
    assert exc.value.reason == reason

def test_enterprise_bundle_equal_to_json():
    parsed = ingest_uploads(uploads(payload())); original = load_sample_evidence()
    assert parsed.bundle.raid == original.raid
    assert parsed.bundle.milestone_plan == original.milestone_plan
    assert parsed.bundle.quality == original.quality
    assert all(m.valid for m in parsed.mappings)
    location = parsed.locations['raid_log.json#/items/0']
    assert location.row == 2 and location.worksheet == 'RAID'
    assert location.columns['title'] == 'Description'

@pytest.mark.parametrize('field', ['raid_log', 'milestone_plan', 'quality_summary'])
def test_csv_matches_xlsx(field):
    files = payload(); name, data = files[field]; sheet, rows = read_xlsx(data, field)
    output = StringIO(); writer = csv.writer(output)
    for number, values in rows:
        writer.writerow([value.date().isoformat() if isinstance(value, datetime) else value for value in values])
    files[field] = (name.replace('.xlsx', '.csv'), output.getvalue().encode())
    result = ingest_uploads(uploads(files)); original=load_sample_evidence()
    assert result.bundle.raid == original.raid
    assert result.bundle.milestone_plan == original.milestone_plan
    assert result.bundle.quality == original.quality
    assert any(loc.worksheet is None and loc.row == 2 for loc in result.locations.values())

@pytest.mark.parametrize('key,value,expected', [
    ('due_date','9 September 2026',date(2026,9,9)), ('forecast','TBC',None),
    ('baseline',datetime(2026,9,9),date(2026,9,9)), ('critical','YES',True),
    ('critical','No',False), ('critical',1,True), ('critical',0,False),
    ('sev2_open','11',11), ('affects','Testing; UAT',['Testing','UAT']),
])
def test_conversions(key, value, expected):
    assert convert(key, value) == expected

@pytest.mark.parametrize('value', ['09/10/2026', '9 September', '9099-09-09'])
def test_ambiguous_dates_rejected(value):
    with pytest.raises(ValueError): convert('due_date', value)

def test_blanks_trimmed_aliases_row_numbers_and_missing_fields():
    source = b'  Milestone  ,Planned Date,Revised Date,On Critical Path,Completed\n\nCertification,2026-09-09,TBC,yes,no\n'
    sheet, rows = read_csv(source, 'milestone_plan')
    result, mapping, locations = map_table('plan.csv', 'milestone_plan', sheet, rows, load_sample_evidence().weekly_status)
    assert result.milestones[0].id == 'ROW-3' and locations[0].row == 3
    assert mapping.mapped_fields['Milestone'] == 'title'
    with pytest.raises(IngestionError) as exc:
        map_table('plan.csv','milestone_plan',None,[(1,['Planned Date']),(2,['2026-09-09'])],load_sample_evidence().weekly_status)
    assert 'title' in exc.value.mapping.missing_required_fields and not exc.value.mapping.valid

def test_ambiguous_alias_columns_rejected():
    with pytest.raises(IngestionError) as exc:
        map_table('plan.csv','milestone_plan',None,[(1,['Milestone','Task']),(2,['A','B'])],load_sample_evidence().weekly_status)
    assert exc.value.reason == 'ambiguous_columns'

def test_malformed_xlsx_and_formula_rejected():
    with pytest.raises(IngestionError) as exc: read_xlsx(b'not a workbook','raid_log')
    assert exc.value.reason == 'malformed'
    book=Workbook(); book.active.append(['ID']); book.active.append(['=1+1']); out=BytesIO(); book.save(out)
    with pytest.raises(IngestionError) as exc: read_xlsx(out.getvalue(),'raid_log')
    assert exc.value.reason == 'formulas'

def test_first_worksheet_and_blank_rows():
    book=Workbook(); sheet=book.active; sheet.title='RAID'; sheet.append(['ID','Type']); sheet.append([None,None]); sheet.append(['R1','risk'])
    book.create_sheet('Ignored').append(['=1+1']); out=BytesIO(); book.save(out)
    name,rows=read_xlsx(out.getvalue(),'raid_log')
    assert name=='RAID' and [number for number,_ in rows]==[1,3]

def test_cross_format_assessments_and_provenance(monkeypatch, extraction_result):
    mock=Mock(return_value=extraction_result); monkeypatch.setattr(assessment_orchestrator,'extract_status_report',mock)
    client=TestClient(app); legacy=payload(False); legacy['status_report']=('report.txt',legacy['status_report'][1])
    a=client.post(PATH,files=legacy); b=client.post(PATH,files=payload())
    assert a.status_code == b.status_code == 200, (a.text,b.text)
    a,b=a.json(),b.json()
    for result in (a,b):
        assert result['reported_health']=='GREEN' and result['programpulse_health']=='AMBER' and result['status_conflict'] is True
        assert next(d for d in result['dimension_summary'] if d['dimension']=='DEPENDENCIES')['status']=='RED'
    assert [(d['dimension'],d['status']) for d in a['dimension_summary']] == [(d['dimension'],d['status']) for d in b['dimension_summary']]
    assert [c['reason'] for c in a['top_governance_challenges']] == [c['reason'] for c in b['top_governance_challenges']]
    assert [c['description'] for c in a['top_recommended_actions']] == [c['description'] for c in b['top_recommended_actions']]
    assert any('Sheet: RAID | Row: 2' in e['source_reference'] for e in b['key_evidence'])
    assert mock.call_count == 2

@pytest.mark.parametrize('name,data,reason', [('bad.xlsm',b'x','extension'),('bad.xlsx',b'x','malformed'),('large.xlsx',b'x'*1_000_001,'size')])
def test_bad_enterprise_uploads_never_extract(monkeypatch, extraction_result, name, data, reason):
    mock=Mock(return_value=extraction_result); monkeypatch.setattr(assessment_orchestrator,'extract_status_report',mock)
    files=payload(); files['raid_log']=(name,data)
    response=TestClient(app).post(PATH,files=files)
    assert response.status_code==422 and response.json()['detail']['reason']==reason
    mock.assert_not_called()
