from io import BytesIO
from openpyxl import load_workbook
from .models import IngestionError
from .safety import check_zip, MAX_ROWS, MAX_COLUMNS


def read_xlsx(data: bytes, field: str) -> tuple[str, list[tuple[int, list]]]:
    check_zip(data, field)
    book = None
    try:
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        sheet.reset_dimensions()  # Do not trust the workbook's declared used range.
        rows = []
        for index, row in enumerate(sheet.iter_rows(), 1):
            if index > MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                raise IngestionError(field, "resource_limit")
            if any(cell.data_type == "f" for cell in row):
                raise IngestionError(field, "formulas")
            values = [cell.value.strip() if isinstance(cell.value, str) else cell.value for cell in row]
            if any(value is not None and value != "" for value in values):
                rows.append((index, values))
        return sheet.title, rows
    except IngestionError:
        raise
    except Exception:
        raise IngestionError(field, "malformed") from None
    finally:
        if book:
            book.close()
